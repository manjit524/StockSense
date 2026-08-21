"""
================================================================================
 REPORTS_DATA.PY — full growth/technical report: how much did SYMBOL move,
 and what do the standard technical indicators say about it
================================================================================
Uses yfinance historical downloads directly (not the local backfilled CSVs)
because reports need arbitrary date ranges uniformly across NSE (.NS), BSE
(.BO), and US tickers.

HONESTY NOTES — read before presenting this anywhere official:
- News Sentiment uses yfinance's own `.news` feed + a small built-in
  positive/negative keyword list. It is a simple heuristic, not a trained
  NLP/sentiment model. Coverage varies a lot by ticker — many symbols will
  have zero recent articles, in which case it reports "Neutral" by default.
- AI Score / AI Trend / AI Recommendation / Confidence are a transparent,
  rule-based composite of the RSI/MACD/SMA/return signals computed in this
  same report. They are NOT the trained GRU model's output — the GRU predicts
  one price point ahead and isn't built to score a historical window. Keep
  that distinction clear in your project writeup.
- Sharpe Ratio uses a flat assumed risk-free rate (6.5% INR, 4.5% USD) —
  a simplification, not a live rate. Worth stating as an assumption.
- Beta is computed against ^NSEI (Nifty 50) for NSE/BSE tickers and ^GSPC
  (S&P 500) for US tickers. Not computed for crypto/commodities — there's no
  single obvious benchmark for those, so it's left as "—".
================================================================================
"""

import numpy as np
import pandas as pd
import yfinance as yf
from datetime import datetime, timedelta

PERIOD_OPTIONS = [
    {"value": "7", "label": "1 Week"},
    {"value": "14", "label": "2 Weeks"},
    {"value": "30", "label": "1 Month"},
    {"value": "90", "label": "3 Months"},
    {"value": "180", "label": "6 Months"},
    {"value": "365", "label": "1 Year"},
]

_PERIOD_LABELS = {int(p["value"]): p["label"] for p in PERIOD_OPTIONS}

POS_WORDS = {
    "beat", "beats", "surge", "surges", "rally", "rallies", "gain", "gains", "up", "rise", "rises",
    "growth", "strong", "record", "upgrade", "upgraded", "outperform", "bullish", "profit", "profits",
    "positive", "boost", "jump", "jumps", "soar", "soars", "higher", "wins", "win",
}
NEG_WORDS = {
    "miss", "misses", "fall", "falls", "drop", "drops", "plunge", "plunges", "decline", "declines",
    "weak", "downgrade", "downgraded", "underperform", "bearish", "loss", "losses", "negative",
    "cut", "cuts", "lower", "crash", "crashes", "slump", "slumps", "fraud", "lawsuit", "probe", "warns",
}


def _to_yf_ticker(symbol, exchange):
    if exchange == "NSE":
        return f"{symbol}.NS"
    if exchange == "BSE":
        return f"{symbol}.BO"
    return symbol  # US, or already a raw yfinance ticker (crypto, commodities, etc.)


def _benchmark_ticker(exchange):
    if exchange in ("NSE", "BSE"):
        return "^NSEI"
    if exchange == "US":
        return "^GSPC"
    return None


def _period_label(period_days):
    return _PERIOD_LABELS.get(int(period_days), f"{period_days} Days")


def _flatten(df):
    df = df.reset_index()
    df.columns = [c[0] if isinstance(c, tuple) else c for c in df.columns]
    return df.sort_values("Date").reset_index(drop=True)


def _rsi(series, period=14):
    delta = series.diff()
    gain = delta.clip(lower=0)
    loss = -delta.clip(upper=0)
    avg_gain = gain.rolling(period).mean()
    avg_loss = loss.rolling(period).mean()
    rs = avg_gain / avg_loss
    return 100 - (100 / (1 + rs))


def _macd(series, fast=12, slow=26):
    ema_fast = series.ewm(span=fast, adjust=False).mean()
    ema_slow = series.ewm(span=slow, adjust=False).mean()
    return ema_fast - ema_slow


def _max_drawdown_pct(series):
    cummax = series.cummax()
    drawdown = (series - cummax) / cummax
    return float(drawdown.min() * 100)


def _simple_sentiment(headlines):
    if not headlines:
        return 0.0, "Neutral", 0
    score_total, counted = 0, 0
    for h in headlines:
        words = set(h.lower().replace(",", " ").replace(".", " ").split())
        pos, neg = len(words & POS_WORDS), len(words & NEG_WORDS)
        if pos or neg:
            score_total += (pos - neg)
            counted += 1
    if counted == 0:
        return 0.0, "Neutral", len(headlines)
    avg = score_total / counted
    score = max(-1.0, min(1.0, avg / 3))
    label = "Positive" if score > 0.15 else "Negative" if score < -0.15 else "Neutral"
    return score, label, len(headlines)


def generate_full_report(symbol, exchange, period_days):
    """Returns a dict with every field needed by both the short and long
    report column sets, or None if price history isn't available."""
    yf_ticker = _to_yf_ticker(symbol, exchange)
    period_days = int(period_days)

    end_date = datetime.today()
    lookback_buffer = max(100, period_days + 60)  # room for stable SMA50/MACD/RSI
    fetch_start = end_date - timedelta(days=period_days + lookback_buffer)

    try:
        hist = yf.download(yf_ticker, start=fetch_start.strftime("%Y-%m-%d"),
                            end=end_date.strftime("%Y-%m-%d"), progress=False)
    except Exception:
        hist = None
    if hist is None or hist.empty:
        return None
    hist = _flatten(hist)

    period_start_date = pd.Timestamp(end_date - timedelta(days=period_days))
    period_df = hist[hist["Date"] >= period_start_date]
    if len(period_df) < 2:
        period_df = hist.tail(max(2, min(len(hist), period_days)))
    if len(period_df) < 2:
        return None

    open_price = float(period_df.iloc[0]["Open"])
    high_price = float(period_df["High"].max())
    low_price = float(period_df["Low"].min())
    close_price = float(period_df.iloc[-1]["Close"])
    start_price = float(period_df.iloc[0]["Close"])
    price_change = close_price - start_price
    return_pct = (price_change / start_price * 100) if start_price else 0.0
    avg_price = float(period_df["Close"].mean())

    has_volume = "Volume" in period_df.columns
    volume_latest = float(period_df.iloc[-1]["Volume"]) if has_volume else None
    avg_volume = float(period_df["Volume"].mean()) if has_volume else None

    prev_period_start = period_start_date - timedelta(days=period_days)
    prev_period_df = hist[(hist["Date"] >= prev_period_start) & (hist["Date"] < period_start_date)]
    volume_change_pct = None
    if avg_volume and has_volume and not prev_period_df.empty:
        prev_avg_vol = float(prev_period_df["Volume"].mean())
        if prev_avg_vol:
            volume_change_pct = (avg_volume - prev_avg_vol) / prev_avg_vol * 100

    hist["rsi"] = _rsi(hist["Close"])
    hist["macd"] = _macd(hist["Close"])
    hist["sma20"] = hist["Close"].rolling(20).mean()
    hist["sma50"] = hist["Close"].rolling(50).mean()

    last_row = hist[hist["Date"] <= period_df.iloc[-1]["Date"]].iloc[-1]
    rsi_value = None if pd.isna(last_row["rsi"]) else float(last_row["rsi"])
    macd_value = None if pd.isna(last_row["macd"]) else float(last_row["macd"])
    sma20_value = None if pd.isna(last_row["sma20"]) else float(last_row["sma20"])
    sma50_value = None if pd.isna(last_row["sma50"]) else float(last_row["sma50"])

    period_returns = period_df["Close"].pct_change().dropna()
    volatility = float(period_returns.std() * (252 ** 0.5) * 100) if len(period_returns) > 1 else None
    max_drawdown = _max_drawdown_pct(period_df["Close"])

    beta_value = None
    benchmark_ticker = _benchmark_ticker(exchange)
    if benchmark_ticker and len(period_returns) > 1:
        try:
            bench_hist = _flatten(yf.download(
                benchmark_ticker, start=period_df.iloc[0]["Date"].strftime("%Y-%m-%d"),
                end=end_date.strftime("%Y-%m-%d"), progress=False,
            ))
            bench_returns = bench_hist["Close"].pct_change().dropna()
            n = min(len(bench_returns), len(period_returns))
            if n > 1:
                cov = np.cov(period_returns.values[-n:], bench_returns.values[-n:])[0][1]
                var = np.var(bench_returns.values[-n:])
                beta_value = float(cov / var) if var else None
        except Exception:
            beta_value = None

    risk_free_rate = 0.065 if exchange in ("NSE", "BSE") else 0.045  # assumption, not live
    sharpe_ratio = None
    if len(period_returns) > 1 and period_returns.std():
        annual_return = period_returns.mean() * 252
        annual_vol = period_returns.std() * (252 ** 0.5)
        sharpe_ratio = float((annual_return - risk_free_rate) / annual_vol)

    sentiment_score, news_sentiment, headline_count = 0.0, "Neutral", 0
    try:
        news_items = yf.Ticker(yf_ticker).news or []
        headlines = [item.get("title", "") for item in news_items if item.get("title")]
        sentiment_score, news_sentiment, headline_count = _simple_sentiment(headlines)
    except Exception:
        pass

    # Rule-based composite — see the honesty note at the top of this file.
    signals = []
    if rsi_value is not None:
        if 45 <= rsi_value <= 65:
            signals.append(0.6)
        elif rsi_value > 70:
            signals.append(-0.4)
        elif rsi_value < 30:
            signals.append(-0.2)
        else:
            signals.append(0.2)
    if macd_value is not None:
        signals.append(0.6 if macd_value > 0 else -0.6)
    if sma20_value and sma50_value:
        signals.append(0.5 if sma20_value > sma50_value else -0.5)
    if sma20_value:
        signals.append(0.4 if close_price > sma20_value else -0.4)
    signals.append(max(-1.0, min(1.0, return_pct / 10)))

    raw_score = sum(signals) / len(signals) if signals else 0.0
    ai_score = round((raw_score + 1) / 2 * 100, 1)
    ai_trend = "Bullish" if ai_score > 60 else "Bearish" if ai_score < 40 else "Neutral"
    ai_recommendation = "Buy" if ai_score > 65 else "Sell" if ai_score < 35 else "Hold"
    agreement = (sum(1 for s in signals if (s > 0) == (raw_score > 0)) / len(signals)) if signals else 0.0
    confidence = round(agreement * 100, 1)

    currency = "$" if exchange == "US" else "₹"

    def money(v): return f"{currency}{v:,.2f}" if v is not None else "—"
    def pct(v): return f"{v:+.2f}%" if v is not None else "—"
    def num(v, d=2): return f"{v:,.{d}f}" if v is not None else "—"

    # Range position calculations (where current close sits between period low and high)
    range_span = (high_price - low_price) if (high_price - low_price) > 0 else 1.0
    range_pos_pct = max(0.0, min(100.0, ((close_price - low_price) / range_span) * 100))

    # Moving average comparisons
    sma20_diff_pct = ((close_price - sma20_value) / sma20_value * 100) if sma20_value else None
    sma50_diff_pct = ((close_price - sma50_value) / sma50_value * 100) if sma50_value else None

    # Build per-day OHLCV rows for the price-history table.
    has_adj = "Adj Close" in period_df.columns
    chart_points = []
    prev_close = None
    for _, row in period_df.iterrows():
        adj_c = float(row["Adj Close"]) if has_adj and not pd.isna(row.get("Adj Close")) else float(row["Close"])
        c_val = float(row["Close"])
        o_val = float(row["Open"])
        h_val = float(row["High"])
        l_val = float(row["Low"])
        
        # Daily return based on prior day's close (or open if first day)
        base_ref = prev_close if prev_close is not None else o_val
        d_chg = ((c_val - base_ref) / base_ref * 100) if base_ref else 0.0
        prev_close = c_val

        chart_points.append({
            "date":             row["Date"].strftime("%Y-%m-%d"),
            "price":            c_val,   # kept for the SVG sparkline
            "open":             o_val,
            "high":             h_val,
            "low":              l_val,
            "close":            c_val,
            "daily_change_pct": d_chg,
            "adj_close":        adj_c,
            "volume":           int(row["Volume"]) if has_volume and not pd.isna(row.get("Volume")) else None,
        })

    # Contextual Takeaway
    dir_word = "appreciated" if price_change >= 0 else "retreated"
    period_str = _period_label(period_days)
    rsi_context = "in overbought territory (>70)" if (rsi_value and rsi_value > 70) else ("in oversold territory (<30)" if (rsi_value and rsi_value < 30) else f"in a balanced momentum zone ({rsi_value:.1f})" if rsi_value else "neutral")
    
    takeaway = (
        f"{symbol} has {dir_word} {abs(return_pct):.2f}% over the {period_str} timeframe, trading between "
        f"{currency}{low_price:,.2f} and {currency}{high_price:,.2f}. Technical momentum is {rsi_context}. "
        f"The composite quantitative engine registers a {ai_recommendation.upper()} stance ({ai_trend} Trend) with {confidence:.1f}% signal alignment."
    )

    return {
        "symbol": symbol, "exchange": exchange, "currency": currency,
        "period_label": _period_label(period_days), "period_days": period_days,
        "start_date": period_df.iloc[0]["Date"].strftime("%Y-%m-%d"),
        "end_date": period_df.iloc[-1]["Date"].strftime("%Y-%m-%d"),
        "direction": "up" if price_change >= 0 else "down",
        "chart_points": chart_points,
        "start_price": start_price, "end_price": close_price,
        "abs_change": price_change, "pct_change": return_pct,

        "open": open_price, "high": high_price, "low": low_price, "close": close_price,
        "range_span": high_price - low_price, "range_pos_pct": range_pos_pct,
        "price_change": price_change, "return_pct": return_pct, "avg_price": avg_price,
        "volume": volume_latest, "avg_volume": avg_volume, "volume_change_pct": volume_change_pct,
        "rsi": rsi_value, "macd": macd_value, "sma20": sma20_value, "sma50": sma50_value,
        "sma20_diff_pct": sma20_diff_pct, "sma50_diff_pct": sma50_diff_pct,
        "volatility": volatility, "max_drawdown": max_drawdown,
        "beta": beta_value, "sharpe_ratio": sharpe_ratio,
        "news_sentiment": news_sentiment, "sentiment_score": sentiment_score,
        "headline_count": headline_count,
        "ai_score": ai_score, "ai_trend": ai_trend,
        "ai_recommendation": ai_recommendation, "confidence": confidence,
        "takeaway": takeaway,

        "open_fmt": money(open_price), "high_fmt": money(high_price),
        "low_fmt": money(low_price), "close_fmt": money(close_price),
        "avg_price_fmt": money(avg_price), "price_change_fmt": money(price_change),
        "return_pct_fmt": pct(return_pct),
        "volume_fmt": f"{volume_latest:,.0f}" if volume_latest is not None else "—",
        "avg_volume_fmt": f"{avg_volume:,.0f}" if avg_volume is not None else "—",
        "volume_change_pct_fmt": pct(volume_change_pct),
        "rsi_fmt": num(rsi_value, 1), "macd_fmt": num(macd_value, 3),
        "sma20_fmt": money(sma20_value), "sma50_fmt": money(sma50_value),
        "volatility_fmt": f"{volatility:.2f}%" if volatility is not None else "—",
        "max_drawdown_fmt": f"{max_drawdown:.2f}%" if max_drawdown is not None else "—",
        "beta_fmt": num(beta_value, 2), "sharpe_fmt": num(sharpe_ratio, 2),
        "sentiment_score_fmt": num(sentiment_score, 2),
        "ai_score_fmt": f"{ai_score:.1f}", "confidence_fmt": f"{confidence:.1f}%",
    }


# Backward-compatible name used before this file grew — points to the same function.
generate_growth_report = generate_full_report


SHORT_COLUMNS = [
    ("Stock", "symbol"), ("Period", "period_label"), ("Open", "open_fmt"), ("High", "high_fmt"),
    ("Low", "low_fmt"), ("Close", "close_fmt"), ("Price Change", "price_change_fmt"),
    ("Return %", "return_pct_fmt"), ("Volume", "volume_fmt"), ("Avg Volume", "avg_volume_fmt"),
    ("RSI", "rsi_fmt"), ("MACD", "macd_fmt"), ("Volatility", "volatility_fmt"),
    ("Max Drawdown", "max_drawdown_fmt"), ("AI Score", "ai_score_fmt"), ("AI Trend", "ai_trend"),
    ("AI Recommendation", "ai_recommendation"),
]

LONG_COLUMNS = [
    ("Stock", "symbol"), ("Period", "period_label"), ("Open", "open_fmt"), ("High", "high_fmt"),
    ("Low", "low_fmt"), ("Close", "close_fmt"), ("Price Change", "price_change_fmt"),
    ("Return %", "return_pct_fmt"), ("Average Price", "avg_price_fmt"),
    ("Volume", "volume_fmt"), ("Avg Volume", "avg_volume_fmt"), ("Volume Change %", "volume_change_pct_fmt"),
    ("RSI", "rsi_fmt"), ("SMA 20", "sma20_fmt"), ("SMA 50", "sma50_fmt"), ("MACD", "macd_fmt"),
    ("Volatility", "volatility_fmt"), ("Max Drawdown", "max_drawdown_fmt"),
    ("Beta", "beta_fmt"), ("Sharpe Ratio", "sharpe_fmt"),
    ("News Sentiment", "news_sentiment"), ("Sentiment Score", "sentiment_score_fmt"),
    ("AI Score", "ai_score_fmt"), ("AI Trend", "ai_trend"),
    ("AI Recommendation", "ai_recommendation"), ("Confidence", "confidence_fmt"),
]


def get_report_columns(period_days):
    """1 week / 2 weeks get the short set; everything longer gets the full set."""
    return SHORT_COLUMNS if int(period_days) <= 14 else LONG_COLUMNS


# ==============================================================================
# EXPORTS — PDF and Excel versions of the same report data
# ==============================================================================

def build_pdf_report(report):
    import io
    from reportlab.lib.pagesizes import A4
    from reportlab.lib import colors
    from reportlab.platypus import SimpleDocTemplate, Table, TableStyle, Paragraph, Spacer
    from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle

    buffer = io.BytesIO()
    doc = SimpleDocTemplate(buffer, pagesize=A4, rightMargin=36, leftMargin=36, topMargin=36, bottomMargin=36)
    styles = getSampleStyleSheet()

    # Custom styles
    title_style = ParagraphStyle(
        'DocTitle',
        parent=styles['Title'],
        fontName='Helvetica-Bold',
        fontSize=18,
        leading=22,
        textColor=colors.HexColor('#0F172A'),
        alignment=0
    )
    subtitle_style = ParagraphStyle(
        'DocSubtitle',
        parent=styles['Normal'],
        fontName='Helvetica',
        fontSize=10,
        leading=14,
        textColor=colors.HexColor('#64748B'),
        alignment=0
    )
    section_style = ParagraphStyle(
        'DocSection',
        parent=styles['Heading2'],
        fontName='Helvetica-Bold',
        fontSize=12,
        leading=16,
        textColor=colors.HexColor('#0F172A'),
        spaceBefore=12,
        spaceAfter=6
    )
    takeaway_style = ParagraphStyle(
        'DocTakeaway',
        parent=styles['Normal'],
        fontName='Helvetica-Oblique',
        fontSize=9.5,
        leading=14,
        textColor=colors.HexColor('#334155'),
        spaceBefore=4,
        spaceAfter=10
    )
    disclaimer_style = ParagraphStyle(
        'DocDisclaimer',
        parent=styles['Italic'],
        fontName='Helvetica-Oblique',
        fontSize=7.5,
        leading=10,
        textColor=colors.HexColor('#94A3B8'),
    )

    elements = []

    # Header Title
    elements.append(Paragraph(f"StockSense Quantitative Performance Report", title_style))
    elements.append(Paragraph(f"Symbol: <b>{report['symbol']}</b> ({report['exchange']}) | Period: <b>{report['period_label']}</b> ({report['start_date']} to {report['end_date']})", subtitle_style))
    elements.append(Spacer(1, 10))

    # Executive Narrative Takeaway
    if report.get("takeaway"):
        elements.append(Paragraph(f"<b>Executive Summary:</b> {report['takeaway']}", takeaway_style))

    # Top KPI Table
    cur = report["currency"]
    dir_sym = "▲ +" if report["direction"] == "up" else "▼ "
    kpi_data = [
        ["Total Return", "Starting Price", "Closing Price", "Period Range", "AI Signal"],
        [
            f"{dir_sym}{report['pct_change']:.2f}%",
            f"{cur}{report['start_price']:,.2f}",
            f"{cur}{report['end_price']:,.2f}",
            f"{cur}{report['low']:,.2f} - {cur}{report['high']:,.2f}",
            f"{report['ai_recommendation']} ({report['ai_trend']})"
        ]
    ]
    kpi_table = Table(kpi_data, colWidths=[100, 105, 105, 120, 95])
    kpi_table.setStyle(TableStyle([
        ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#F1F5F9")),
        ("TEXTCOLOR", (0, 0), (-1, 0), colors.HexColor("#475569")),
        ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
        ("FONTSIZE", (0, 0), (-1, 0), 8),
        ("BOTTOMPADDING", (0, 0), (-1, 0), 5),
        ("BACKGROUND", (0, 1), (-1, 1), colors.HexColor("#F8FAFC")),
        ("TEXTCOLOR", (0, 1), (-1, 1), colors.HexColor("#0F172A")),
        ("FONTNAME", (0, 1), (-1, 1), "Helvetica-Bold"),
        ("FONTSIZE", (0, 1), (-1, 1), 9.5),
        ("ALIGN", (0, 0), (-1, -1), "CENTER"),
        ("GRID", (0, 0), (-1, -1), 0.5, colors.HexColor("#CBD5E1")),
    ]))
    elements.append(kpi_table)
    elements.append(Spacer(1, 14))

    # Metrics Breakdown Table
    elements.append(Paragraph("Technical & Quantitative Metrics Breakdown", section_style))
    columns = get_report_columns(report["period_days"])
    metric_rows = [["Metric Parameter", "Observed Value"]] + [[label, str(report.get(key, "—"))] for label, key in columns[2:]]
    metrics_table = Table(metric_rows, colWidths=[240, 285])
    metrics_table.setStyle(TableStyle([
        ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#0F172A")),
        ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
        ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
        ("FONTSIZE", (0, 0), (-1, -1), 8.5),
        ("GRID", (0, 0), (-1, -1), 0.5, colors.HexColor("#E2E8F0")),
        ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.HexColor("#F8FAFC")]),
        ("TOPPADDING", (0, 0), (-1, -1), 4),
        ("BOTTOMPADDING", (0, 0), (-1, -1), 4),
    ]))
    elements.append(metrics_table)
    elements.append(Spacer(1, 14))

    # Price History Table (Last 15 days in PDF to keep page clean)
    elements.append(Paragraph("Recent OHLCV Price Sequence", section_style))
    price_rows = [["Date", "Open", "High", "Low", "Close", "Daily %", "Volume"]]
    pts_to_show = report["chart_points"][-15:] if len(report["chart_points"]) > 15 else report["chart_points"]
    for pt in pts_to_show:
        vol_str = f"{pt['volume']:,}" if pt.get("volume") is not None else "—"
        d_pct_str = f"{pt.get('daily_change_pct', 0.0):+.2f}%"
        price_rows.append([
            pt["date"],
            f"{cur}{pt['open']:.2f}",
            f"{cur}{pt['high']:.2f}",
            f"{cur}{pt['low']:.2f}",
            f"{cur}{pt['close']:.2f}",
            d_pct_str,
            vol_str,
        ])
    price_table = Table(price_rows, colWidths=[75, 75, 75, 75, 75, 75, 75])
    price_table.setStyle(TableStyle([
        ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#1E293B")),
        ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
        ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
        ("FONTSIZE", (0, 0), (-1, -1), 8),
        ("ALIGN", (1, 0), (-1, -1), "RIGHT"),
        ("GRID", (0, 0), (-1, -1), 0.5, colors.HexColor("#E2E8F0")),
        ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.HexColor("#F8FAFC")]),
        ("TOPPADDING", (0, 0), (-1, -1), 3),
        ("BOTTOMPADDING", (0, 0), (-1, -1), 3),
    ]))
    elements.append(price_table)
    elements.append(Spacer(1, 14))

    # Disclaimer Note
    elements.append(Paragraph(
        "Disclaimer: Generated for informational & analytical purposes only. StockSense AI Scores and signals are rule-based composites "
        "and do not constitute SEBI/SEC registered financial advice. Past performance is no guarantee of future returns.",
        disclaimer_style,
    ))

    doc.build(elements)
    buffer.seek(0)
    return buffer


def build_excel_report(report):
    import io
    import openpyxl
    from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
    from openpyxl.utils import get_column_letter

    wb = openpyxl.Workbook()

    # Sheet 1: Executive Summary & Technical Metrics
    ws1 = wb.active
    ws1.title = "Executive Summary"
    ws1.views.sheetView[0].showGridLines = True

    # Styling palettes
    navy_fill = PatternFill(start_color="0F172A", end_color="0F172A", fill_type="solid")
    header_fill = PatternFill(start_color="1E293B", end_color="1E293B", fill_type="solid")
    sub_fill = PatternFill(start_color="F1F5F9", end_color="F1F5F9", fill_type="solid")
    alt_fill = PatternFill(start_color="F8FAFC", end_color="F8FAFC", fill_type="solid")
    white_fill = PatternFill(start_color="FFFFFF", end_color="FFFFFF", fill_type="solid")

    green_fill = PatternFill(start_color="ECFDF5", end_color="ECFDF5", fill_type="solid")
    red_fill = PatternFill(start_color="FEF2F2", end_color="FEF2F2", fill_type="solid")

    font_title = Font(name="Calibri", size=14, bold=True, color="FFFFFF")
    font_header = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
    font_bold = Font(name="Calibri", size=10, bold=True, color="0F172A")
    font_regular = Font(name="Calibri", size=10, color="334155")
    font_green = Font(name="Calibri", size=10, bold=True, color="059669")
    font_red = Font(name="Calibri", size=10, bold=True, color="DC2626")

    thin_border = Border(
        left=Side(style='thin', color="E2E8F0"),
        right=Side(style='thin', color="E2E8F0"),
        top=Side(style='thin', color="E2E8F0"),
        bottom=Side(style='thin', color="E2E8F0")
    )

    # 1. Title Banner
    ws1.merge_cells("A1:D1")
    title_cell = ws1["A1"]
    title_cell.value = f"STOCKSENSE QUANTITATIVE PERFORMANCE & GROWTH REPORT"
    title_cell.font = font_title
    title_cell.fill = navy_fill
    title_cell.alignment = Alignment(horizontal="center", vertical="center")
    ws1.row_dimensions[1].height = 36

    # 2. Metadata Subtitle
    ws1.merge_cells("A2:D2")
    sub_cell = ws1["A2"]
    sub_cell.value = f"Asset: {report['symbol']} ({report['exchange']}) | Horizon: {report['period_label']} ({report['start_date']} to {report['end_date']}) | Currency: {report['currency']}"
    sub_cell.font = Font(name="Calibri", size=9.5, italic=True, color="64748B")
    sub_cell.alignment = Alignment(horizontal="center", vertical="center")
    ws1.row_dimensions[2].height = 20

    # 3. KPI Highlights Cards
    kpis = [
        ("TOTAL RETURN", f"{report['pct_change']:+.2f}%", font_green if report['direction'] == 'up' else font_red, green_fill if report['direction'] == 'up' else red_fill),
        ("START PRICE", f"{report['currency']}{report['start_price']:,.2f}", font_bold, sub_fill),
        ("CLOSE PRICE", f"{report['currency']}{report['end_price']:,.2f}", font_bold, sub_fill),
        ("AI RECOMMENDATION", f"{report['ai_recommendation']} ({report['ai_trend']})", font_bold, sub_fill),
    ]

    for col_idx, (kpi_lbl, kpi_val, kpi_font, kpi_bg) in enumerate(kpis, start=1):
        c1 = ws1.cell(row=4, column=col_idx, value=kpi_lbl)
        c1.font = Font(name="Calibri", size=8.5, bold=True, color="64748B")
        c1.alignment = Alignment(horizontal="center", vertical="center")
        c1.fill = sub_fill
        c1.border = thin_border

        c2 = ws1.cell(row=5, column=col_idx, value=kpi_val)
        c2.font = kpi_font
        c2.alignment = Alignment(horizontal="center", vertical="center")
        c2.fill = kpi_bg
        c2.border = thin_border
    ws1.row_dimensions[4].height = 18
    ws1.row_dimensions[5].height = 24

    # 4. Executive Narrative Takeaway
    ws1.merge_cells("A7:D7")
    narrative_cell = ws1["A7"]
    narrative_cell.value = f"Summary: {report.get('takeaway', '')}"
    narrative_cell.font = Font(name="Calibri", size=9.5, italic=True, color="334155")
    narrative_cell.alignment = Alignment(horizontal="left", vertical="center", wrap_text=True)
    ws1.row_dimensions[7].height = 36

    # 5. Full Metrics Table
    columns = get_report_columns(report["period_days"])
    ws1.cell(row=9, column=1, value="Category").fill = header_fill
    ws1.cell(row=9, column=1).font = font_header
    ws1.cell(row=9, column=2, value="Metric Parameter").fill = header_fill
    ws1.cell(row=9, column=2).font = font_header
    ws1.cell(row=9, column=3, value="Observed Value").fill = header_fill
    ws1.cell(row=9, column=3).font = font_header
    ws1.cell(row=9, column=4, value="Signal / Context").fill = header_fill
    ws1.cell(row=9, column=4).font = font_header
    ws1.row_dimensions[9].height = 22

    row_idx = 10
    for label, key in columns:
        if key in ["symbol", "period_label"]:
            continue
        
        # Categorize
        if key in ["open_fmt", "high_fmt", "low_fmt", "close_fmt", "price_change_fmt", "return_pct_fmt", "avg_price_fmt", "volatility_fmt", "max_drawdown_fmt", "beta_fmt"]:
            cat = "Price & Volatility"
        elif key in ["rsi_fmt", "macd_fmt", "sma20_fmt", "sma50_fmt", "volume_fmt", "avg_volume_fmt", "volume_change_pct_fmt"]:
            cat = "Technical Indicators"
        else:
            cat = "AI & Sentiment Ratios"

        val_str = str(report.get(key, "—"))
        
        # Context interpretation
        ctx = "—"
        if key == "rsi_fmt" and report.get("rsi"):
            ctx = "Overbought (>70)" if report["rsi"] > 70 else ("Oversold (<30)" if report["rsi"] < 30 else "Neutral Momentum")
        elif key == "macd_fmt" and report.get("macd") is not None:
            ctx = "Bullish Momentum" if report["macd"] > 0 else "Bearish Momentum"
        elif key == "ai_trend":
            ctx = f"{report.get('confidence', 70):.1f}% Model Confidence"
        elif key == "volatility_fmt":
            ctx = "Annualized Volatility"
        elif key == "max_drawdown_fmt":
            ctx = "Peak-to-Trough Decline"

        bg = alt_fill if row_idx % 2 == 0 else white_fill
        for col_i, text in enumerate([cat, label, val_str, ctx], start=1):
            cell = ws1.cell(row=row_idx, column=col_i, value=text)
            cell.font = font_regular
            cell.fill = bg
            cell.border = thin_border
            if col_i == 3:
                cell.font = font_bold
        ws1.row_dimensions[row_idx].height = 20
        row_idx += 1

    # Adjust column widths for Sheet 1
    ws1.column_dimensions['A'].width = 24
    ws1.column_dimensions['B'].width = 26
    ws1.column_dimensions['C'].width = 20
    ws1.column_dimensions['D'].width = 32

    # Sheet 2: Price History (OHLCV)
    ws2 = wb.create_sheet(title="Price History")
    ws2.views.sheetView[0].showGridLines = True

    p_headers = ["Date", f"Open ({report['currency']})", f"High ({report['currency']})", f"Low ({report['currency']})", f"Close ({report['currency']})", "Daily Return (%)", f"Adj Close ({report['currency']})", "Volume"]
    for col_i, h in enumerate(p_headers, start=1):
        cell = ws2.cell(row=1, column=col_i, value=h)
        cell.font = font_header
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal="center", vertical="center")
        cell.border = thin_border
    ws2.row_dimensions[1].height = 24

    for r_i, pt in enumerate(report["chart_points"], start=2):
        bg = alt_fill if r_i % 2 == 0 else white_fill
        d_chg = pt.get("daily_change_pct", 0.0)
        chg_font = font_green if d_chg >= 0 else font_red

        row_vals = [
            pt["date"],
            pt["open"],
            pt["high"],
            pt["low"],
            pt["close"],
            d_chg,
            pt["adj_close"],
            pt["volume"] if pt["volume"] is not None else "—"
        ]

        for col_i, val in enumerate(row_vals, start=1):
            cell = ws2.cell(row=r_i, column=col_i, value=val)
            cell.font = chg_font if col_i == 6 else font_regular
            cell.fill = bg
            cell.border = thin_border
            if col_i in [2, 3, 4, 5, 7] and isinstance(val, (int, float)):
                cell.number_format = '#,##0.00'
            elif col_i == 6 and isinstance(val, (int, float)):
                cell.number_format = '+0.00%;-0.00%;0.00%'
            elif col_i == 8 and isinstance(val, (int, float)):
                cell.number_format = '#,##0'
        ws2.row_dimensions[r_i].height = 19

    # Auto-fit columns for Sheet 2
    for col in ws2.columns:
        max_len = max(len(str(cell.value or '')) for cell in col)
        col_letter = get_column_letter(col[0].column)
        ws2.column_dimensions[col_letter].width = max(max_len + 4, 14)

    buffer = io.BytesIO()
    wb.save(buffer)
    buffer.seek(0)
    return buffer
